from __future__ import annotations

import json
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import pandas as pd
from openpyxl import load_workbook

from app.ingestion.models import ColumnProfile, SourceFileProfile
from app.config import get_settings

ALLOWED_EXTENSIONS = {".csv", ".xlsx"}
IDENTIFIER_HINTS = {"employee_id", "emp_id", "worker_id", "staff_id", "email"}
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
PHONE_RE = re.compile(r"\+?\d[\d\s().-]{6,}\d")
DATE_PATTERNS = {
    r"^\d{4}-\d{2}-\d{2}$": "YYYY-MM-DD",
    r"^\d{2}/\d{2}/\d{4}$": "DD/MM/YYYY_OR_MM/DD/YYYY",
    r"^\d{2}-[A-Za-z]{3}-\d{4}$": "DD-MMM-YYYY",
}


class IngestionError(ValueError):
    pass


def sanitize_filename(name: str) -> str:
    clean = Path(name).name
    clean = unicodedata.normalize("NFKC", clean)
    clean = re.sub(r"[^A-Za-z0-9._-]+", "_", clean).strip("._")
    if not clean:
        raise IngestionError("Filename is empty after sanitization")
    return clean[:180]


def _json_value(value: Any) -> Any:
    if value is None or (not isinstance(value, (list, dict)) and pd.isna(value)):
        return None
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return value.isoformat()
    if hasattr(value, "item"):
        return value.item()
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _read_csv(path: Path) -> tuple[pd.DataFrame, str, None]:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            frame = pd.read_csv(path, encoding=encoding, on_bad_lines="error", dtype=object, nrows=get_settings().max_source_rows + 1)
            _check_shape(frame)
            frame.columns = _unique_headers(list(frame.columns))
            return frame, encoding, None
        except UnicodeDecodeError as exc:
            last_error = exc
        except (pd.errors.ParserError, pd.errors.EmptyDataError) as exc:
            raise IngestionError(f"Malformed CSV: {exc}") from exc
    raise IngestionError(f"Unable to decode CSV: {last_error}")


def _read_xlsx(path: Path) -> tuple[pd.DataFrame, None, str]:
    try:
        with ZipFile(path) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise IngestionError("Expanded Excel file exceeds 50 MB. Split it into smaller files.")
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        sheet = workbook[workbook.sheetnames[0]]
        rows: list[list[Any]] = []
        for row in sheet.iter_rows():
            if len(rows) > get_settings().max_source_rows or len(row) > get_settings().max_source_columns:
                workbook.close()
                raise IngestionError("File exceeds 50,000 rows or 100 columns. Split it into smaller files.")
            rows.append([None if cell.data_type == "f" else cell.value for cell in row])
        workbook.close()
    except Exception as exc:
        raise IngestionError(f"Invalid XLSX file: {exc}") from exc
    if not rows:
        raise IngestionError("Spreadsheet is empty")
    raw_headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    if not any(raw_headers):
        raise IngestionError("Spreadsheet header is empty")
    headers = _unique_headers(rows[0])
    return pd.DataFrame(rows[1:], columns=headers), None, sheet.title


def _check_shape(frame: pd.DataFrame) -> None:
    settings = get_settings()
    if len(frame) > settings.max_source_rows or len(frame.columns) > settings.max_source_columns:
        raise IngestionError(f"File exceeds {settings.max_source_rows:,} rows or {settings.max_source_columns} columns. Split it into smaller files.")


def _unique_headers(values: list[Any]) -> list[str]:
    headers: list[str] = []
    seen: set[str] = set()
    occurrences: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        base = str(value).strip() if value is not None else ""
        if not base:
            base = f"unnamed_column_{index}"
        occurrence = occurrences.get(base, 0) + 1
        occurrences[base] = occurrence
        candidate = base if occurrence == 1 else f"{base}__{occurrence}"
        while candidate in seen:
            occurrence += 1
            occurrences[base] = occurrence
            candidate = f"{base}__{occurrence}"
        seen.add(candidate)
        headers.append(candidate)
    return headers


def read_source(path: Path) -> tuple[pd.DataFrame, str | None, str | None]:
    if path.suffix.lower() == ".csv":
        return _read_csv(path)
    if path.suffix.lower() == ".xlsx":
        return _read_xlsx(path)
    raise IngestionError("Only .csv and .xlsx files are supported")


def _kind(series: pd.Series) -> str:
    values = series.dropna()
    if values.empty:
        return "unknown"
    strings = values.astype(str).str.strip()
    if strings.str.fullmatch(r"[-+]?\d+").mean() >= 0.9:
        return "integer"
    if strings.str.fullmatch(r"[-+]?(?:\d+\.?\d*|\.\d+)").mean() >= 0.9:
        return "number"
    if strings.map(lambda item: any(re.match(pattern, item) for pattern in DATE_PATTERNS)).mean() >= 0.7:
        return "date"
    if strings.map(lambda item: bool(EMAIL_RE.match(item))).mean() >= 0.7:
        return "email"
    return "string"


def _mask(value: Any, column_name: str) -> Any:
    safe = _json_value(value)
    if safe is None:
        return None
    text = str(safe)
    normalized = re.sub(r"\W+", "_", column_name.lower()).strip("_")
    if "email" in normalized or EMAIL_RE.match(text):
        local, _, domain = text.partition("@")
        return f"{local[:1]}***@{domain}" if domain else "***"
    if "phone" in normalized or "mobile" in normalized or PHONE_RE.fullmatch(text):
        digits = re.sub(r"\D", "", text)
        return f"***{digits[-4:]}" if digits else "***"
    if normalized in IDENTIFIER_HINTS or normalized.endswith("_id"):
        return f"***{text[-3:]}" if text else "***"
    return text[:80]


def profile_frame(frame: pd.DataFrame, file_name: str, encoding: str | None, sheet: str | None) -> SourceFileProfile:
    if frame.empty or len(frame.columns) == 0:
        raise IngestionError("Source file contains no data rows")
    frame = frame.map(_json_value)
    profiles: list[ColumnProfile] = []
    for column_index, raw_name in enumerate(frame.columns):
        name = str(raw_name).strip()
        series = frame.iloc[:, column_index]
        non_null = series.dropna()
        patterns = sorted(
            {("UNAMBIGUOUS_SLASH_DATE" if label == "DD/MM/YYYY_OR_MM/DD/YYYY" and
              (int(value.split('/')[0]) > 12 or int(value.split('/')[1]) > 12 or value.split('/')[0] == value.split('/')[1])
              else label) for value in non_null.astype(str) for pattern, label in DATE_PATTERNS.items() if re.match(pattern, value)}
        )
        normalized = re.sub(r"\W+", "_", name.lower()).strip("_")
        profiles.append(
            ColumnProfile(
                name=name,
                inferred_type=_kind(series),
                null_ratio=round(float(series.isna().mean()), 4),
                unique_ratio=round(float(non_null.nunique(dropna=True) / max(len(non_null), 1)), 4),
                masked_samples=[_mask(value, name) for value in non_null.head(3).tolist()],
                likely_identifier=(normalized in IDENTIFIER_HINTS or normalized.endswith("_id") or (len(non_null) > 0 and non_null.nunique() == len(non_null))),
                date_patterns=patterns,
            )
        )
    return SourceFileProfile(
        file_name=file_name,
        sheet_name=sheet,
        row_count=len(frame),
        columns=profiles,
        duplicate_row_count=int(frame.astype(str).duplicated().sum()),
        encoding=encoding,
    )


def profile_file(path: Path, safe_name: str) -> SourceFileProfile:
    frame, encoding, sheet = read_source(path)
    return profile_frame(frame, safe_name, encoding, sheet)


def profile_to_json(profile: SourceFileProfile) -> str:
    return json.dumps(profile.model_dump(mode="json"), separators=(",", ":"))
